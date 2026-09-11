"""Read-only final-billing evidence. Never mutates reports or close status.

Export rows are evidence of charges, NOT proof that a billing period is closed.
Adapters must obtain period dates and closed status from separate official evidence.
"""
from __future__ import annotations

import calendar
import hashlib
import io
import json
import re
from collections import Counter
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

import openpyxl


HEADERS = {
    "BRL": {
        "charge_id": "Número da tarifa", "charged_on": "Data da tarifa",
        "description": "Detalhe", "amount": "Valor da tarifa",
        "order_id": "Número da venda", "reversed_charge_id": "Tarifa cancelada",
        "status": "Status da tarifa",
    },
    "MXN": {
        "charge_id": "Número del cargo", "charged_on": "Fecha del cargo",
        "description": "Detalle", "amount": "Valor del cargo",
        "order_id": "Número de venta", "reversed_charge_id": "Cargo que bonifica",
        "status": "Estado del cargo",
    },
}


def month_bounds(month: str) -> tuple[date, date]:
    if not isinstance(month, str) or len(month) != 7:
        raise ValueError("月份必须为 YYYY-MM")
    start = date.fromisoformat(month + "-01")
    return start, start.replace(day=calendar.monthrange(start.year, start.month)[1])


def _identifier(value, *, required=False) -> str:
    if isinstance(value, str):
        value = value.strip()
    if value in (None, ""):
        if required:
            raise ValueError("费用编号为空，无法去重")
        return ""
    # Large numeric Excel identifiers may already have lost precision.
    if isinstance(value, bool) or isinstance(value, (float, int)) and (
        abs(value) >= 10**15 or int(value) != value
    ):
        raise ValueError("编号可能丢失精度，请提供保留文本编号的官方文件")
    return str(int(value)) if isinstance(value, (float, int)) else str(value).strip()


def _amount(value) -> Decimal:
    if value is None or isinstance(value, bool):
        raise ValueError("费用金额缺失或格式错误，不能当作0")
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("费用金额格式错误") from exc
    if not result.is_finite():
        raise ValueError("费用金额不是有限数值")
    return result


def _date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    # Never guess whether 08/09 means 8 September or 9 August.
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    raise ValueError("费用日期不明确，不能猜测日/月顺序")


def read_local_export(content: bytes, seller_id: str, currency: str) -> dict:
    """Read BR/MX REPORT by exact headers; preserve signs and reversal references."""
    expected = {"2378517428": "BRL", "3383185411": "MXN"}
    if expected.get(str(seller_id)) != currency:
        raise ValueError("店铺与币种不匹配或尚未配置")
    digest = hashlib.sha256(content).hexdigest()
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if "REPORT" not in workbook.sheetnames:
            raise ValueError("官方账单缺少 REPORT 工作表")
        sheet = workbook["REPORT"]
        headers = [str(c.value or "").strip() for c in sheet[8]]
        indexes = {}
        for key, label in HEADERS[currency].items():
            if headers.count(label) != 1:
                raise ValueError(f"账单字段缺失或重复：{label}")
            indexes[key] = headers.index(label)
        rows = []
        seen = {}
        duplicates = 0
        for number, values in enumerate(sheet.iter_rows(min_row=9, values_only=True), 9):
            if all(v in (None, "") for v in values):
                continue
            raw = {key: values[i] for key, i in indexes.items()}
            try:
                row = {
                    "seller_id": str(seller_id), "currency": currency,
                    "charge_id": _identifier(raw["charge_id"], required=True),
                    "charged_on": _date(raw["charged_on"]).isoformat(),
                    "description": str(raw["description"] or "").strip(),
                    "amount": str(_amount(raw["amount"]).normalize()),
                    "order_id": _identifier(raw["order_id"]),
                    "reversed_charge_id": _identifier(raw["reversed_charge_id"]),
                    "status": str(raw["status"] or "").strip(),
                }
                if not row["description"]:
                    raise ValueError("费用类型为空")
            except ValueError as exc:
                raise ValueError(f"REPORT 第{number}行：{exc}") from exc
            key = row["charge_id"]
            if key in seen:
                if seen[key] != row:
                    raise ValueError(f"同一费用编号内容冲突：{key}")
                duplicates += 1
                continue
            seen[key] = row.copy()
            rows.append({**row, "source_sha256": digest, "source_row": number})
        if not rows:
            raise ValueError("账单无费用记录；零费用须另提供官方证明")
        return {
            "seller_id": str(seller_id), "currency": currency,
            "source_sha256": digest, "rows": rows, "duplicates": duplicates,
            "month_counts": dict(sorted(Counter(r["charged_on"][:7] for r in rows).items())),
            "closed_period_verified": False,
            "ready_for_final_confirmation": False,
        }
    finally:
        workbook.close()


def check_closed_coverage(month: str, periods: list[dict], seller_id: str, currency: str) -> dict:
    """Validate normalized official-period evidence, not charge min/max dates.

    Callers must fetch trusted official metadata; this function is not an API
    accepting an operator-supplied CLOSED flag as proof.
    """
    start, end = month_bounds(month)
    covered = set()
    blockers = []
    for period in periods:
        if str(period.get("seller_id")) != str(seller_id) or period.get("currency") != currency:
            blockers.append("账期店铺或币种不匹配")
            continue
        if not re.fullmatch(r"[0-9a-f]{64}", str(period.get("source_sha256", ""))) or not period.get("source_reference"):
            blockers.append("账期缺少官方来源证据")
            continue
        left, right = _date(period.get("date_from")), _date(period.get("date_to"))
        if right < left:
            raise ValueError("账期结束早于开始")
        if right < start or left > end:
            continue
        if period.get("status") != "CLOSED":
            blockers.append("覆盖目标月份的账期尚未正式结束")
            continue
        if period.get("documents_complete") is not True:
            blockers.append("账期发票、贷项通知单或Full明细尚未核实齐全")
            continue
        day = max(start, left)
        while day <= min(end, right):
            covered.add(day)
            day += timedelta(days=1)
    missing = [ (start + timedelta(days=i)).isoformat()
                for i in range((end-start).days+1)
                if start + timedelta(days=i) not in covered ]
    if missing:
        blockers.append("官方已结束账期未覆盖整月")
    return {"complete": not blockers, "missing_dates": missing, "blockers": sorted(set(blockers))}


def compare_charge_sources(automatic: list[dict], official: list[dict]) -> dict:
    """Exact raw-charge A/B check before any profit allocation.

    Inputs must be independently fetched: A from the API, B from the official
    export. An empty comparison never passes. This does NOT authorize close.
    """
    fields = ("charged_on", "description", "amount", "order_id", "reversed_charge_id", "status")

    def index(rows):
        result = {}
        for row in rows:
            if not re.fullmatch(r"[0-9a-f]{64}", str(row.get("source_sha256", ""))):
                raise ValueError("对账行缺少来源文件指纹")
            key = tuple(_identifier(row.get(k), required=True) for k in ("seller_id", "currency", "charge_id"))
            if any(k not in row for k in fields):
                raise ValueError("对账字段不完整")
            normalized = {k: str(row[k]).strip() for k in fields}
            normalized["charged_on"] = _date(row["charged_on"]).isoformat()
            normalized["amount"] = str(_amount(row["amount"]).normalize())
            if key in result and result[key][0] != normalized:
                raise ValueError("多个来源中的同一费用编号内容冲突")
            result[key] = (normalized, row["source_sha256"])
        return result

    a, b = index(automatic), index(official)
    differences = []
    for key in sorted(a.keys() | b.keys()):
        if key not in a or key not in b:
            differences.append({"key": list(key), "kind": "missing_in_automatic" if key not in a else "missing_in_official"})
            continue
        changed = {field: {"automatic": a[key][0][field], "official": b[key][0][field]}
                   for field in fields if a[key][0][field] != b[key][0][field]}
        if changed:
            differences.append({"key": list(key), "kind": "field_difference", "fields": changed,
                                "automatic_source": a[key][1], "official_source": b[key][1]})
    evidence = {"automatic": [(list(k), *a[k]) for k in sorted(a)],
                "official": [(list(k), *b[k]) for k in sorted(b)]}
    fingerprint = hashlib.sha256(json.dumps(evidence, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return {"matched": bool(a and b) and not differences, "automatic_count": len(a),
            "official_count": len(b), "differences": differences, "evidence_sha256": fingerprint,
            "ready_for_final_confirmation": False}
