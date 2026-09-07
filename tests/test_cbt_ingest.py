import io
import unittest
from unittest.mock import AsyncMock, patch

import openpyxl
from fastapi import HTTPException

from app import cbt_ingest, main


class _Response:
    def __init__(self, payload, status_code=200):
        self._payload = payload
        self.status_code = status_code

    def json(self):
        return self._payload


class _WriteClient:
    def __init__(self):
        self.calls = []
        self.created_records = []

    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return False

    async def post(self, url, **kwargs):
        self.calls.append(("POST", url, kwargs.get("json")))
        if url.endswith("/batch_create"):
            self.created_records = kwargs["json"]["records"]
            return _Response({"code": 0, "data": {"records": [
                {"record_id": f"new-{index}"}
                for index, _record in enumerate(self.created_records, start=1)
            ]}})
        if url.endswith("/batch_delete"):
            return _Response({"code": 0})
        if url.endswith("/batch_update"):
            return _Response({"code": 0})
        raise AssertionError(f"unexpected POST {url}")

    async def put(self, url, **kwargs):
        raise AssertionError(f"in-place update is unsafe: {url}")


class _ReadFailureClient:
    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return False

    async def get(self, url, **kwargs):
        return _Response({"code": 99999, "msg": "failed"}, status_code=500)


def _xlsx_bytes(workbook):
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _orders_file(month, sku):
    import calendar

    year, month_number = (int(part) for part in month.split("-"))
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Orders US"
    sheet.cell(6, 3, "Order date")
    for row_number, day in enumerate((1, 15, calendar.monthrange(year, month_number)[1]), start=7):
        row = [None] * 73
        row[2] = __import__("datetime").date(year, month_number, day).strftime(
            "%B %d, %Y 10:00 AM GMT-06:00"
        )
        row[9] = 1
        row[10] = 10
        row[11] = -1
        row[13] = 0
        row[14] = -2
        row[16] = -0.5
        row[17] = 0
        row[18] = 6.5
        row[24] = sku
        row[25] = f"LIST-{sku}"
        row[26] = f"Title {sku}"
        sheet.append(row)
    return _xlsx_bytes(workbook)


def _ads_file(month, sku):
    year, month_number = (int(part) for part in month.split("-"))
    import calendar
    import datetime

    start = datetime.date(year, month_number, 1)
    end = datetime.date(year, month_number, calendar.monthrange(year, month_number)[1])
    workbook = openpyxl.Workbook()
    help_sheet = workbook.active
    help_sheet.title = "Help"
    help_sheet.cell(11, 7, "Time Period")
    help_sheet.cell(11, 12, f"{start:%d %B %Y} - {end:%d %B %Y}")
    report = workbook.create_sheet("Report by ads")
    report.cell(2, 1, "From")
    report.cell(2, 2, "To")
    report.cell(3, 1, start.strftime("%d-%b-%Y"))
    report.cell(3, 2, end.strftime("%d-%b-%Y"))
    report.cell(3, 5, f"LIST-{sku}")
    report.cell(3, 13, 2.5)
    return _xlsx_bytes(workbook)


def _bill_file(month):
    year, month_number = (int(part) for part in month.split("-"))
    import calendar
    import datetime

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "REPORT"
    sheet.cell(8, 2, "Fee release date")
    sheet.cell(9, 2, datetime.datetime(year, month_number, 1))
    sheet.cell(9, 4, "Storage service fee")
    sheet.cell(9, 8, 3)
    sheet.cell(10, 2, datetime.datetime(year, month_number, calendar.monthrange(year, month_number)[1]))
    sheet.cell(10, 4, "Service fee")
    sheet.cell(10, 8, 1)
    return _xlsx_bytes(workbook)


def _file(name, token, modified_time):
    return {"name": name, "token": token, "type": "file", "modified_time": modified_time}


class CbtIngestPeriodSafetyTests(unittest.IsolatedAsyncioTestCase):
    async def test_bitable_read_failure_is_not_treated_as_an_empty_table(self):
        with patch.object(cbt_ingest.httpx, "AsyncClient", return_value=_ReadFailureClient()):
            with self.assertRaisesRegex(RuntimeError, "报表读取失败"):
                await cbt_ingest._bitable_all("token")

    async def test_run_selects_files_whose_contents_match_requested_month(self):
        files = [
            _file("latest Orders.xlsx", "orders-aug", "20"),
            _file("older Orders.xlsx", "orders-jul", "10"),
            _file("latest report-pads.xlsx", "ads-aug", "20"),
            _file("older report-pads.xlsx", "ads-jul", "10"),
            _file("latest BILL.xlsx", "bill-aug", "20"),
            _file("older BILL.xlsx", "bill-jul", "10"),
        ]
        payloads = {
            "orders-aug": _orders_file("2026-08", "AUG-SKU"),
            "orders-jul": _orders_file("2026-07", "JUL-SKU"),
            "ads-aug": _ads_file("2026-08", "AUG-SKU"),
            "ads-jul": _ads_file("2026-07", "JUL-SKU"),
            "bill-aug": _bill_file("2026-08"),
            "bill-jul": _bill_file("2026-07"),
        }

        async def download(_token, file_token):
            return payloads[file_token]

        with (
            patch.object(cbt_ingest, "_ft", AsyncMock(return_value="token")),
            patch.object(cbt_ingest, "_list_folder", AsyncMock(return_value=files)),
            patch.object(cbt_ingest, "_download", side_effect=download),
            patch.object(cbt_ingest.lingxing, "fetch_all_products", AsyncMock(return_value={})),
        ):
            result = await cbt_ingest.run("2026-07", commit=False, folder_token="folder")

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["files"]["orders"], "older Orders.xlsx")
        self.assertEqual(result["files"]["ads"], "older report-pads.xlsx")
        self.assertEqual(result["files"]["bill"], "older BILL.xlsx")
        self.assertEqual(result["sku_count"], 1)
        self.assertEqual(result["totals"]["units"], 3)
        self.assertEqual(result["totals"]["K"], 30)

    async def test_orders_parser_includes_first_data_row_and_excludes_prior_month_boundary(self):
        workbook = openpyxl.load_workbook(io.BytesIO(_orders_file("2026-08", "AUG-SKU")))
        sheet = workbook.active
        row = [None] * 73
        row[2] = "July 31, 2026 11:59 PM GMT-06:00"
        row[9] = 9
        row[10] = 99
        row[11] = -9
        row[13] = 0
        row[14] = -9
        row[16] = -9
        row[17] = 0
        row[18] = 72
        row[24] = "BOUNDARY-SKU"
        row[25] = "LIST-BOUNDARY"
        row[26] = "Prior month boundary"
        sheet.append(row)
        data = _xlsx_bytes(workbook)

        rows, _listing_map = cbt_ingest._parse_orders(data, "2026-08")

        self.assertEqual(set(rows), {"AUG-SKU"})
        self.assertEqual(rows["AUG-SKU"]["orders"], 3)
        self.assertEqual(rows["AUG-SKU"]["units"], 3)
        self.assertEqual(rows["AUG-SKU"]["K"], 30)
        self.assertNotIn("BOUNDARY-SKU", rows)

    async def test_run_rejects_month_when_any_required_export_does_not_match(self):
        files = [
            _file("Orders.xlsx", "orders-aug", "20"),
            _file("report-pads.xlsx", "ads-aug", "20"),
            _file("BILL.xlsx", "bill-aug", "20"),
        ]
        payloads = {
            "orders-aug": _orders_file("2026-08", "AUG-SKU"),
            "ads-aug": _ads_file("2026-08", "AUG-SKU"),
            "bill-aug": _bill_file("2026-08"),
        }

        async def download(_token, file_token):
            return payloads[file_token]

        with (
            patch.object(cbt_ingest, "_ft", AsyncMock(return_value="token")),
            patch.object(cbt_ingest, "_list_folder", AsyncMock(return_value=files)),
            patch.object(cbt_ingest, "_download", side_effect=download),
            patch.object(cbt_ingest.lingxing, "fetch_all_products", AsyncMock(return_value={})),
        ):
            result = await cbt_ingest.run("2026-07", commit=False, folder_token="folder")

        self.assertEqual(result["status"], "error")
        self.assertIn("2026-07", result["msg"])

    async def test_endpoint_turns_validation_error_into_non_2xx_failure(self):
        with patch.object(cbt_ingest, "run", AsyncMock(return_value={"status": "error", "msg": "wrong month"})):
            with self.assertRaises(HTTPException) as raised:
                await main.cbt_ingest(month="2026-08", commit=True)

        self.assertEqual(raised.exception.status_code, 422)
        self.assertIn("wrong month", str(raised.exception.detail))

    async def test_commit_archives_entire_old_scope_and_preserves_stale_skus(self):
        files = [
            _file("Orders.xlsx", "orders", "20"),
            _file("report-pads.xlsx", "ads", "20"),
            _file("BILL.xlsx", "bill", "20"),
        ]
        payloads = {
            "orders": _orders_file("2026-08", "AUG-SKU"),
            "ads": _ads_file("2026-08", "AUG-SKU"),
            "bill": _bill_file("2026-08"),
        }
        existing = [
            {"record_id": "old-current", "fields": {"店铺": "ML CBT-FULL (1502236229)", "周期": "month_2026-08", "SKU": "AUG-SKU", "头程成本(RMB)": 11, "海外仓成本(RMB)": 12}},
            {"record_id": "old-stale", "fields": {"店铺": "ML CBT-FULL (1502236229)", "周期": "month_2026-08", "SKU": "STALE-SKU"}},
        ]
        client = _WriteClient()
        reads = 0

        async def download(_token, file_token):
            return payloads[file_token]

        async def bitable_all(_token):
            nonlocal reads
            reads += 1
            if reads == 1:
                return existing
            fresh = [{"record_id": "new-1", "fields": client.created_records[0]["fields"]}]
            if reads == 2:
                return existing + fresh
            archived = [
                {**item, "fields": {**item["fields"], "周期": "month_2026-08_original_20260907"}}
                for item in existing
            ]
            return archived + fresh

        with (
            patch.object(cbt_ingest, "_ft", AsyncMock(return_value="token")),
            patch.object(cbt_ingest, "_list_folder", AsyncMock(return_value=files)),
            patch.object(cbt_ingest, "_download", side_effect=download),
            patch.object(cbt_ingest, "_ensure_full_field", AsyncMock()),
            patch.object(cbt_ingest, "_bitable_all", side_effect=bitable_all),
            patch.object(cbt_ingest.lingxing, "fetch_all_products", AsyncMock(return_value={})),
            patch.object(cbt_ingest.httpx, "AsyncClient", return_value=client),
        ):
            result = await cbt_ingest.run(
                "2026-08",
                commit=True,
                folder_token="folder",
                preserve_existing_as="month_2026-08_original_20260907",
            )

        delete_calls = [call for call in client.calls if call[1].endswith("/batch_delete")]
        self.assertEqual(delete_calls, [])
        update_calls = [call for call in client.calls if call[1].endswith("/batch_update")]
        self.assertEqual(
            update_calls[0][2]["records"],
            [
                {"record_id": "old-current", "fields": {"周期": "month_2026-08_original_20260907"}},
                {"record_id": "old-stale", "fields": {"周期": "month_2026-08_original_20260907"}},
            ],
        )
        self.assertEqual(client.created_records[0]["fields"]["头程成本(RMB)"], 11)
        self.assertEqual(client.created_records[0]["fields"]["海外仓成本(RMB)"], 12)
        self.assertEqual(result["rows_archived"], 2)
        self.assertEqual(result["stale_rows_archived"], 1)
        self.assertEqual(result["rows_verified"], 1)

    async def test_commit_blocks_different_existing_scope_without_backup_label(self):
        files = [
            _file("Orders.xlsx", "orders", "20"),
            _file("report-pads.xlsx", "ads", "20"),
            _file("BILL.xlsx", "bill", "20"),
        ]
        payloads = {
            "orders": _orders_file("2026-08", "AUG-SKU"),
            "ads": _ads_file("2026-08", "AUG-SKU"),
            "bill": _bill_file("2026-08"),
        }
        existing = [{
            "record_id": "old-1",
            "fields": {"店铺": "ML CBT-FULL (1502236229)", "周期": "month_2026-08", "SKU": "AUG-SKU"},
        }]

        async def download(_token, file_token):
            return payloads[file_token]

        with (
            patch.object(cbt_ingest, "_ft", AsyncMock(return_value="token")),
            patch.object(cbt_ingest, "_list_folder", AsyncMock(return_value=files)),
            patch.object(cbt_ingest, "_download", side_effect=download),
            patch.object(cbt_ingest, "_ensure_full_field", AsyncMock()),
            patch.object(cbt_ingest, "_bitable_all", AsyncMock(return_value=existing)),
            patch.object(cbt_ingest.lingxing, "fetch_all_products", AsyncMock(return_value={})),
        ):
            with self.assertRaisesRegex(RuntimeError, "preserve_existing_as"):
                await cbt_ingest.run("2026-08", commit=True, folder_token="folder")

    async def test_commit_rolls_back_new_rows_when_pre_delete_totals_do_not_match(self):
        files = [
            _file("Orders.xlsx", "orders", "20"),
            _file("report-pads.xlsx", "ads", "20"),
            _file("BILL.xlsx", "bill", "20"),
        ]
        payloads = {
            "orders": _orders_file("2026-08", "AUG-SKU"),
            "ads": _ads_file("2026-08", "AUG-SKU"),
            "bill": _bill_file("2026-08"),
        }
        existing = [{"record_id": "old-1", "fields": {"店铺": "ML CBT-FULL (1502236229)", "周期": "month_2026-08", "SKU": "AUG-SKU"}}]
        client = _WriteClient()
        reads = 0

        async def download(_token, file_token):
            return payloads[file_token]

        async def bitable_all(_token):
            nonlocal reads
            reads += 1
            if reads == 1:
                return existing
            wrong_fields = dict(client.created_records[0]["fields"])
            wrong_fields["营收(原币)"] = 999
            return existing + [{"record_id": "new-1", "fields": wrong_fields}]

        with (
            patch.object(cbt_ingest, "_ft", AsyncMock(return_value="token")),
            patch.object(cbt_ingest, "_list_folder", AsyncMock(return_value=files)),
            patch.object(cbt_ingest, "_download", side_effect=download),
            patch.object(cbt_ingest, "_ensure_full_field", AsyncMock()),
            patch.object(cbt_ingest, "_bitable_all", side_effect=bitable_all),
            patch.object(cbt_ingest.lingxing, "fetch_all_products", AsyncMock(return_value={})),
            patch.object(cbt_ingest.httpx, "AsyncClient", return_value=client),
        ):
            with self.assertRaisesRegex(RuntimeError, "写入前核验失败"):
                await cbt_ingest.run(
                    "2026-08",
                    commit=True,
                    folder_token="folder",
                    preserve_existing_as="month_2026-08_original_20260907",
                )

        delete_payloads = [call[2]["records"] for call in client.calls if call[1].endswith("/batch_delete")]
        self.assertEqual(delete_payloads, [["new-1"]])

    async def test_commit_rejects_per_sku_value_swap_even_when_aggregate_matches(self):
        workbook = openpyxl.load_workbook(io.BytesIO(_orders_file("2026-08", "SKU-A")))
        sheet = workbook.active
        row = [None] * 73
        row[2] = "August 20, 2026 10:00 AM GMT-06:00"
        row[9] = 1
        row[10] = 50
        row[11] = -5
        row[13] = 0
        row[14] = -5
        row[16] = -2
        row[17] = 0
        row[18] = 38
        row[24] = "SKU-B"
        row[25] = "LIST-SKU-B"
        row[26] = "Title SKU-B"
        sheet.append(row)
        files = [
            _file("Orders.xlsx", "orders", "20"),
            _file("report-pads.xlsx", "ads", "20"),
            _file("BILL.xlsx", "bill", "20"),
        ]
        payloads = {
            "orders": _xlsx_bytes(workbook),
            "ads": _ads_file("2026-08", "SKU-A"),
            "bill": _bill_file("2026-08"),
        }
        client = _WriteClient()
        reads = 0

        async def download(_token, file_token):
            return payloads[file_token]

        async def bitable_all(_token):
            nonlocal reads
            reads += 1
            if reads == 1:
                return []
            first = dict(client.created_records[0]["fields"])
            second = dict(client.created_records[1]["fields"])
            first["营收(原币)"], second["营收(原币)"] = second["营收(原币)"], first["营收(原币)"]
            return [
                {"record_id": "new-1", "fields": first},
                {"record_id": "new-2", "fields": second},
            ]

        with (
            patch.object(cbt_ingest, "_ft", AsyncMock(return_value="token")),
            patch.object(cbt_ingest, "_list_folder", AsyncMock(return_value=files)),
            patch.object(cbt_ingest, "_download", side_effect=download),
            patch.object(cbt_ingest, "_ensure_full_field", AsyncMock()),
            patch.object(cbt_ingest, "_bitable_all", side_effect=bitable_all),
            patch.object(cbt_ingest.lingxing, "fetch_all_products", AsyncMock(return_value={})),
            patch.object(cbt_ingest.httpx, "AsyncClient", return_value=client),
        ):
            with self.assertRaisesRegex(RuntimeError, "SKU=SKU-A"):
                await cbt_ingest.run("2026-08", commit=True, folder_token="folder")

        delete_payloads = [call[2]["records"] for call in client.calls if call[1].endswith("/batch_delete")]
        self.assertEqual(delete_payloads, [["new-1", "new-2"]])


if __name__ == "__main__":
    unittest.main()
