import tempfile
import unittest
import io
import zipfile
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch
from pathlib import Path

from app.billing_upload import UploadStore
from app.billing_upload_parse import parse


class UploadTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.store = UploadStore(str(Path(self.folder.name) / "intake.sqlite"))
        self.session = self.store.prepare()

    def save(self, data=b"%PDF-1.4\nproof", name="bill.pdf", seller="2378517428", kind="period"):
        return self.store.save(self.session["token"], seller, kind, name, data)

    def test_persistent_receipt_dedup_and_expiring_access(self):
        with tempfile.TemporaryDirectory() as folder:
            path = str(Path(folder) / "intake.sqlite")
            store = UploadStore(path)
            session = store.prepare(now=100)
            receipt = store.save(session["token"], "2378517428", "period", "bill.pdf", b"%PDF-1.4\nproof", now=101)
            reopened = UploadStore(path)
            duplicate = reopened.save(session["token"], "2378517428", "period", "bill.pdf", b"%PDF-1.4\nproof", now=102)
            self.assertEqual(receipt["id"], duplicate["id"])
            self.assertTrue(duplicate["duplicate"])
            status = reopened.status(session["token"], now=103)
            self.assertEqual(len(status["files"]), 1)
            self.assertFalse(status["ready_for_final_confirmation"])
            self.assertEqual(reopened.download(session["token"], receipt["id"], now=104)[1], b"%PDF-1.4\nproof")
            with self.assertRaises(PermissionError):
                reopened.status(session["token"], now=86501)

    def test_invalid_scope_signature_empty_and_traversal_rejected(self):
        for kwargs in ({"seller":"wrong"}, {"kind":"approve"}, {"data":b""}, {"data":b"<script>"}, {"name":"../bill.pdf"}, {"name":"bill.exe"}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                self.save(**kwargs)
        with self.assertRaises(PermissionError):
            self.store.save("wrong", "2378517428", "period", "bill.pdf", b"%PDF-1.4")
        self.assertEqual(self.store.status(self.session["token"])["files"], [])

    def test_same_name_different_content_preserved_and_concurrent_duplicate(self):
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda _: self.save(), range(4)))
        self.assertEqual(len({r["id"] for r in results}), 1)
        self.save(data=b"%PDF-1.4\nchanged")
        files = self.store.status(self.session["token"])["files"]
        self.assertEqual(len(files), 2)
        self.assertTrue(all(f["same_name_conflict"] for f in files))

    def test_revocation_and_cross_batch_download(self):
        receipt = self.save()
        self.store.revoke(self.session["id"])
        with self.assertRaises(PermissionError):
            self.store.status(self.session["token"])
        other = self.store.prepare()
        with self.assertRaises(PermissionError):
            self.store.download(other["token"], receipt["id"])

    def test_single_active_grant_and_registered_card_reused(self):
        self.assertTrue(self.store.prepare()["delivery_uncertain"])
        self.store.register(self.session["id"], "om_test")
        self.assertEqual(self.store.prepare()["message_id"], "om_test")
        with self.assertRaises(ValueError):
            self.store.register(self.session["id"], "om_other")

    def test_quota_duplicate_and_readback_failure(self):
        self.save()
        with patch("app.billing_upload.MAX_BATCH", 20):
            self.assertTrue(self.save()["duplicate"])
            with self.assertRaises(ValueError):
                self.save(data=b"%PDF-1.4\nother")
        with patch("app.billing_upload.MAX_FILE", 4), self.assertRaises(ValueError):
            self.save()

    def test_zip_embedded_program_and_expansion_rejected(self):
        for extra in ("../evil", "xl/vbaProject.bin", "xl/embeddings/oleObject.bin", "xl/externalLinks/link.xml"):
            out = io.BytesIO()
            with zipfile.ZipFile(out, "w") as z:
                z.writestr("[Content_Types].xml", "x")
                z.writestr("xl/workbook.xml", "x")
                z.writestr(extra, "x")
            with self.assertRaises(ValueError):
                self.save(out.getvalue(), "bill.xlsx")

    def test_queue_restart_lease_and_cbt_never_passes(self):
        receipt = self.save(seller="1502236229", kind="charges")
        claimed = self.store.claim(now=100)
        self.assertIsNone(UploadStore(self.store.path).claim(now=101))
        next_claim = self.store.claim(now=191)
        self.assertEqual(claimed[0], next_claim[0])
        self.store.finish(*claimed, {"state":"parsed"})
        self.assertEqual(self.store.status(self.session["token"])["files"][0]["state"], "processing")
        result = parse(self.store.path, receipt["id"])
        self.store.finish(*next_claim, result)
        status = self.store.status(self.session["token"])
        self.assertEqual(status["files"][0]["state"], "needs_review")
        self.assertFalse(status["ready_for_final_confirmation"])

    def test_worker_timeout_terminates_and_preserves_original(self):
        import subprocess
        from app.billing_upload_api import process_one
        receipt = self.save()
        with patch("app.billing_upload_api.subprocess.run", side_effect=subprocess.TimeoutExpired("parse", 60)) as run:
            self.assertTrue(process_one(self.store))
            self.assertEqual(run.call_args.kwargs["timeout"], 60)
        self.assertEqual(self.store.status(self.session["token"])["files"][0]["state"], "failed")
        self.assertTrue(self.store.download(self.session["token"], receipt["id"])[1].startswith(b"%PDF"))

    def test_real_child_process_reads_workbook_and_no_final_pass(self):
        from datetime import datetime
        from openpyxl import Workbook
        from app.billing_upload_api import process_one
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "REPORT"
        for i, label in enumerate(["Número da tarifa", "Data da tarifa", "Detalhe", "Valor da tarifa", "Número da venda", "Tarifa cancelada", "Status da tarifa"], 1):
            sheet.cell(8, i, label)
        for i, value in enumerate(["fee-1", datetime(2026,8,25), "Comissão", -2.5, "order-1", "", "paid"], 1):
            sheet.cell(9, i, value)
        out = io.BytesIO()
        workbook.save(out)
        workbook.close()
        self.store.save(self.session['token'], '2378517428', 'charges', 'synthetic.xlsx', out.getvalue(), source='synthetic')
        self.assertTrue(process_one(self.store))
        status = self.store.status(self.session["token"])
        self.assertEqual(status["files"][0]["state"], "parsed")
        self.assertEqual(status["files"][0]["source"], "synthetic")
        self.assertEqual(status["files"][0]["result"]["rows"], 1)
        self.assertEqual(status["files"][0]["result"]["month_counts"], {"2026-08":1})
        self.assertFalse(status["ready_for_final_confirmation"])


class PortalTests(unittest.TestCase):
    def test_http_upload_download_auth_and_headers(self):
        from fastapi import FastAPI, Header, HTTPException
        from fastapi.testclient import TestClient
        from app.billing_upload_api import install
        with tempfile.TemporaryDirectory() as folder:
            book = UploadStore(str(Path(folder) / "intake.sqlite"))
            app = FastAPI()
            def auth(authorization: str = Header(default="")):
                if authorization != "Bearer internal":
                    raise HTTPException(401)
            install(app, auth)
            with patch("app.billing_upload_api.store", return_value=book):
                client = TestClient(app)
                base = "/report/ml-intake/test"
                self.assertEqual(client.post(base+"/prepare").status_code, 401)
                session = client.post(base+"/prepare", headers={"Authorization":"Bearer internal"}).json()
                headers = {"Authorization":"Bearer "+session["token"]}
                self.assertEqual(client.get(base+"/status").status_code, 401)
                response = client.post(base+"/files", params={"seller":"2378517428","kind":"period","name":"bill.pdf"}, content=b"%PDF-1.4\nproof", headers=headers)
                self.assertEqual(response.status_code, 200)
                fid = response.json()["id"]
                download = client.get(base+"/files/"+fid, headers=headers)
                self.assertEqual(download.content, b"%PDF-1.4\nproof")
                self.assertTrue(download.headers["content-disposition"].startswith("attachment"))
                self.assertEqual(download.headers["cache-control"], "no-store")
                self.assertEqual(download.headers["referrer-policy"], "no-referrer")
                self.assertEqual(client.get(base+"/upload").status_code, 200)
                self.assertEqual(client.get(base+"/portal.js").status_code, 200)
                invalid = io.BytesIO()
                with zipfile.ZipFile(invalid, 'w') as archive:
                    archive.writestr('[Content_Types].xml', '<Types/>')
                    archive.writestr('xl/workbook.xml', '<workbook/>')
                content = bytearray(invalid.getvalue())
                for signature, offset in ((b'PK\x03\x04', 8), (b'PK\x01\x02', 10)):
                    pos = content.find(signature)
                    while pos >= 0:
                        content[pos+offset:pos+offset+2] = (99).to_bytes(2, 'little')
                        pos = content.find(signature, pos+4)
                bad = client.post(base+'/files', params={'seller':'2378517428','kind':'charges','name':'unsupported.xlsx'}, content=bytes(content), headers=headers)
                self.assertEqual(bad.status_code, 400)
                self.assertIn('压缩格式不支持', bad.json()['detail'])
